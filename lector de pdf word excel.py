import tkinter as tk
from tkinter import filedialog
from docx import Document
import openpyxl
import fitz  # Importamos PyMuPDF

def abrir_archivo(tipo, archivo):
    if tipo == "Word":
        document = Document(archivo)
        for paragraph in document.paragraphs:
            print(paragraph.text)
    elif tipo == "Excel":
        workbook = openpyxl.load_workbook(archivo)
        sheet = workbook.active
        for row in sheet.iter_rows(values_only=True):
            print(" ".join(map(str, row)))
    elif tipo == "PDF":
        pdf_document = fitz.open(archivo)
        for page_num in range(len(pdf_document)):
            page = pdf_document[page_num]
            text = page.get_text()
            print(text)
    else:
        print("Tipo de archivo no válido")

def seleccionar_archivo():
    archivo = filedialog.askopenfilename(title="Seleccionar archivo", filetypes=[("Archivos", "*.*")])
    if archivo:
        tipo_archivo = archivo.split(".")[-1]
        if tipo_archivo == "docx":
            abrir_archivo("Word", archivo)
        elif tipo_archivo == "xlsx":
            abrir_archivo("Excel", archivo)
        elif tipo_archivo == "pdf":
            abrir_archivo("PDF", archivo)
        else:
            print("Tipo de archivo no válido")

def main():
    root = tk.Tk()
    root.withdraw()  # Ocultar la ventana principal de tkinter
    seleccionar_archivo()

if __name__ == "__main__":
    main()
